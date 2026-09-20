import os
import io
import shutil
import tempfile
import uuid
from typing import Optional
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response, JSONResponse
import pypdf
from reportlab.lib.pagesizes import letter, landscape, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas
import openpyxl
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

app = FastAPI(title="llosers PDF & Image Platform API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

TEMP_DIR = os.path.join(tempfile.gettempdir(), "lloserr_tmp")
os.makedirs(TEMP_DIR, exist_ok=True)

def cleanup_file(path: str):
    try:
        if os.path.exists(path):
            os.remove(path)
    except Exception as e:
        print(f"Error cleaning up {path}: {e}")

@app.get("/api/health")
def health_check():
    return {
        "status": "healthy",
        "service": "PDF & Image Processing Engine",
        "database": False,
        "ephemeral": True
    }

# -------------------------------------------------------------
# 1. PROTECT PDF (Password Encryption)
# -------------------------------------------------------------
@app.post("/api/protect-pdf")
async def protect_pdf(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    password: str = Form(...)
):
    if not password:
        raise HTTPException(status_code=400, detail="Password is required")
    
    temp_in = os.path.join(TEMP_DIR, f"in_{uuid.uuid4().hex}.pdf")
    temp_out = os.path.join(TEMP_DIR, f"protected_{uuid.uuid4().hex}.pdf")

    try:
        with open(temp_in, "wb") as f:
            content = await file.read()
            f.write(content)

        reader = pypdf.PdfReader(temp_in)
        writer = pypdf.PdfWriter()

        for page in reader.pages:
            writer.add_page(page)

        writer.encrypt(user_password=password, algorithm="AES-256")

        with open(temp_out, "wb") as f:
            writer.write(f)

        background_tasks.add_task(cleanup_file, temp_in)
        background_tasks.add_task(cleanup_file, temp_out)

        return FileResponse(
            temp_out,
            media_type="application/pdf",
            filename=f"protected_{file.filename or 'document.pdf'}"
        )
    except Exception as e:
        cleanup_file(temp_in)
        cleanup_file(temp_out)
        raise HTTPException(status_code=500, detail=f"Failed to protect PDF: {str(e)}")

# -------------------------------------------------------------
# 2. UNLOCK PDF (Remove Password)
# -------------------------------------------------------------
@app.post("/api/unlock-pdf")
async def unlock_pdf(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    password: str = Form(...)
):
    temp_in = os.path.join(TEMP_DIR, f"in_{uuid.uuid4().hex}.pdf")
    temp_out = os.path.join(TEMP_DIR, f"unlocked_{uuid.uuid4().hex}.pdf")

    try:
        with open(temp_in, "wb") as f:
            content = await file.read()
            f.write(content)

        reader = pypdf.PdfReader(temp_in)
        if reader.is_encrypted:
            result = reader.decrypt(password)
            if result == 0:
                raise HTTPException(status_code=400, detail="Incorrect password. Could not unlock PDF.")

        writer = pypdf.PdfWriter()
        for page in reader.pages:
            writer.add_page(page)

        with open(temp_out, "wb") as f:
            writer.write(f)

        background_tasks.add_task(cleanup_file, temp_in)
        background_tasks.add_task(cleanup_file, temp_out)

        return FileResponse(
            temp_out,
            media_type="application/pdf",
            filename=f"unlocked_{file.filename or 'document.pdf'}"
        )
    except HTTPException:
        cleanup_file(temp_in)
        cleanup_file(temp_out)
        raise
    except Exception as e:
        cleanup_file(temp_in)
        cleanup_file(temp_out)
        raise HTTPException(status_code=500, detail=f"Failed to unlock PDF: {str(e)}")

# -------------------------------------------------------------
# 3. WORD TO PDF (.docx -> .pdf)
# -------------------------------------------------------------
@app.post("/api/word-to-pdf")
async def word_to_pdf(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...)
):
    temp_in = os.path.join(TEMP_DIR, f"in_{uuid.uuid4().hex}.docx")
    temp_out = os.path.join(TEMP_DIR, f"out_{uuid.uuid4().hex}.pdf")

    try:
        with open(temp_in, "wb") as f:
            content = await file.read()
            f.write(content)

        doc = Document(temp_in)
        pdf_doc = SimpleDocTemplate(
            temp_out,
            pagesize=letter,
            rightMargin=54, leftMargin=54,
            topMargin=54, bottomMargin=54
        )
        
        styles = getSampleStyleSheet()
        normal_style = styles['Normal']
        normal_style.fontSize = 11
        normal_style.leading = 14
        heading1_style = styles['Heading1']
        heading2_style = styles['Heading2']

        story = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                story.append(Spacer(1, 8))
                continue
            
            style_name = p.style.name.lower()
            if 'heading 1' in style_name:
                story.append(Paragraph(text, heading1_style))
                story.append(Spacer(1, 6))
            elif 'heading 2' in style_name:
                story.append(Paragraph(text, heading2_style))
                story.append(Spacer(1, 4))
            else:
                story.append(Paragraph(text, normal_style))
                story.append(Spacer(1, 4))

        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = [Paragraph(cell.text.strip(), normal_style) for cell in row.cells]
                table_data.append(row_data)
            if table_data:
                col_count = len(table_data[0]) if table_data else 1
                col_width = 504 / max(col_count, 1)
                t = Table(table_data, colWidths=[col_width]*col_count)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#f1f5f9')),
                    ('TEXTCOLOR', (0,0), (-1,-1), colors.HexColor('#0f172a')),
                    ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                    ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
                    ('FONTSIZE', (0,0), (-1,-1), 10),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                    ('TOPPADDING', (0,0), (-1,-1), 6),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
                ]))
                story.append(Spacer(1, 10))
                story.append(t)
                story.append(Spacer(1, 10))

        if not story:
            story.append(Paragraph("Empty Document", normal_style))

        pdf_doc.build(story)

        background_tasks.add_task(cleanup_file, temp_in)
        background_tasks.add_task(cleanup_file, temp_out)

        out_name = (file.filename or "document.docx").rsplit(".", 1)[0] + ".pdf"
        return FileResponse(temp_out, media_type="application/pdf", filename=out_name)
    except Exception as e:
        cleanup_file(temp_in)
        cleanup_file(temp_out)
        raise HTTPException(status_code=500, detail=f"Failed to convert Word to PDF: {str(e)}")

# -------------------------------------------------------------
# 4. EXCEL TO PDF (.xlsx -> .pdf)
# -------------------------------------------------------------
@app.post("/api/excel-to-pdf")
async def excel_to_pdf(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...)
):
    temp_in = os.path.join(TEMP_DIR, f"in_{uuid.uuid4().hex}.xlsx")
    temp_out = os.path.join(TEMP_DIR, f"out_{uuid.uuid4().hex}.pdf")

    try:
        with open(temp_in, "wb") as f:
            content = await file.read()
            f.write(content)

        wb = openpyxl.load_workbook(temp_in, data_only=True)
        pdf_doc = SimpleDocTemplate(
            temp_out,
            pagesize=landscape(A4),
            rightMargin=36, leftMargin=36,
            topMargin=36, bottomMargin=36
        )

        styles = getSampleStyleSheet()
        sheet_title_style = ParagraphStyle(
            'SheetTitle',
            parent=styles['Heading2'],
            textColor=colors.HexColor('#1e293b'),
            fontSize=14,
            leading=18,
            spaceAfter=10
        )
        cell_style = ParagraphStyle(
            'Cell',
            fontName='Helvetica',
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#334155')
        )
        header_cell_style = ParagraphStyle(
            'HeaderCell',
            fontName='Helvetica-Bold',
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#0f172a')
        )

        story = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            story.append(Paragraph(f"Sheet: {sheet_name}", sheet_title_style))
            
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                story.append(Paragraph("<i>(Empty Sheet)</i>", cell_style))
                story.append(Spacer(1, 15))
                continue

            # Limit empty rows/columns
            filtered_rows = []
            for r in rows:
                if any(c is not None and str(c).strip() != "" for c in r):
                    filtered_rows.append(r)

            if not filtered_rows:
                continue

            max_cols = max(len(r) for r in filtered_rows)
            max_cols = min(max_cols, 12)  # fit to page width

            table_data = []
            for idx, r in enumerate(filtered_rows[:100]): # render up to 100 rows per sheet
                row_cells = []
                for c_idx in range(max_cols):
                    val = r[c_idx] if c_idx < len(r) and r[c_idx] is not None else ""
                    style_to_use = header_cell_style if idx == 0 else cell_style
                    row_cells.append(Paragraph(str(val), style_to_use))
                table_data.append(row_cells)

            if table_data:
                page_width = 841.89 - 72 # landscape A4 width minus margins
                col_w = page_width / max_cols
                t = Table(table_data, colWidths=[col_w]*max_cols, repeatRows=1)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e2e8f0')),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                    ('TOPPADDING', (0,0), (-1,-1), 4),
                    ('LEFTPADDING', (0,0), (-1,-1), 4),
                    ('RIGHTPADDING', (0,0), (-1,-1), 4),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
                    ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')])
                ]))
                story.append(t)
                story.append(Spacer(1, 20))

        if not story:
            story.append(Paragraph("Empty Spreadsheet", cell_style))

        pdf_doc.build(story)

        background_tasks.add_task(cleanup_file, temp_in)
        background_tasks.add_task(cleanup_file, temp_out)

        out_name = (file.filename or "spreadsheet.xlsx").rsplit(".", 1)[0] + ".pdf"
        return FileResponse(temp_out, media_type="application/pdf", filename=out_name)
    except Exception as e:
        cleanup_file(temp_in)
        cleanup_file(temp_out)
        raise HTTPException(status_code=500, detail=f"Failed to convert Excel to PDF: {str(e)}")

# -------------------------------------------------------------
# 5. POWERPOINT TO PDF (.pptx -> .pdf)
# -------------------------------------------------------------
@app.post("/api/powerpoint-to-pdf")
async def powerpoint_to_pdf(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...)
):
    temp_in = os.path.join(TEMP_DIR, f"in_{uuid.uuid4().hex}.pptx")
    temp_out = os.path.join(TEMP_DIR, f"out_{uuid.uuid4().hex}.pdf")

    try:
        with open(temp_in, "wb") as f:
            content = await file.read()
            f.write(content)

        prs = Presentation(temp_in)
        pdf_doc = SimpleDocTemplate(
            temp_out,
            pagesize=landscape(letter),
            rightMargin=40, leftMargin=40,
            topMargin=40, bottomMargin=40
        )

        styles = getSampleStyleSheet()
        slide_title_style = ParagraphStyle(
            'SlideTitle',
            parent=styles['Title'],
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#0f172a'),
            alignment=0,
            spaceAfter=16
        )
        bullet_style = ParagraphStyle(
            'SlideBullet',
            parent=styles['Normal'],
            fontSize=13,
            leading=18,
            textColor=colors.HexColor('#334155'),
            spaceAfter=8
        )

        story = []

        for idx, slide in enumerate(prs.slides):
            if idx > 0:
                story.append(PageBreak())

            story.append(Paragraph(f"Slide {idx + 1}", ParagraphStyle('SlideNum', fontSize=9, textColor=colors.HexColor('#94a3b8'), spaceAfter=8)))

            # Extract slide title and content text
            title_text = ""
            body_texts = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            if not title_text and shape == slide.shapes[0]:
                                title_text = text
                            else:
                                body_texts.append(text)

            if title_text:
                story.append(Paragraph(title_text, slide_title_style))
            else:
                story.append(Paragraph(f"Untitled Slide {idx+1}", slide_title_style))

            if body_texts:
                for b_text in body_texts:
                    story.append(Paragraph(f"• {b_text}", bullet_style))
            else:
                story.append(Paragraph("<i>(Visual slide elements)</i>", bullet_style))

        if not story:
            story.append(Paragraph("Empty Presentation", slide_title_style))

        pdf_doc.build(story)

        background_tasks.add_task(cleanup_file, temp_in)
        background_tasks.add_task(cleanup_file, temp_out)

        out_name = (file.filename or "presentation.pptx").rsplit(".", 1)[0] + ".pdf"
        return FileResponse(temp_out, media_type="application/pdf", filename=out_name)
    except Exception as e:
        cleanup_file(temp_in)
        cleanup_file(temp_out)
        raise HTTPException(status_code=500, detail=f"Failed to convert PowerPoint to PDF: {str(e)}")

# -------------------------------------------------------------
# 6. PDF TO WORD (.pdf -> .docx)
# -------------------------------------------------------------
@app.post("/api/pdf-to-word")
async def pdf_to_word(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...)
):
    temp_in = os.path.join(TEMP_DIR, f"in_{uuid.uuid4().hex}.pdf")
    temp_out = os.path.join(TEMP_DIR, f"out_{uuid.uuid4().hex}.docx")

    try:
        with open(temp_in, "wb") as f:
            content = await file.read()
            f.write(content)

        # Try pdf2docx converter if available
        converted = False
        try:
            from pdf2docx import Converter
            cv = Converter(temp_in)
            cv.convert(temp_out, start=0, end=None)
            cv.close()
            converted = True
        except Exception as p2d_err:
            print(f"pdf2docx conversion error: {p2d_err}, falling back to pypdf extractor")

        if not converted or not os.path.exists(temp_out):
            # Pure python fallback using pypdf + docx
            reader = pypdf.PdfReader(temp_in)
            doc = Document()
            for idx, page in enumerate(reader.pages):
                if idx > 0:
                    doc.add_page_break()
                text = page.extract_text()
                if text:
                    for line in text.splitlines():
                        line_s = line.strip()
                        if line_s:
                            doc.add_paragraph(line_s)
                else:
                    doc.add_paragraph(f"[Page {idx+1} Content]")
            doc.save(temp_out)

        background_tasks.add_task(cleanup_file, temp_in)
        background_tasks.add_task(cleanup_file, temp_out)

        out_name = (file.filename or "document.pdf").rsplit(".", 1)[0] + ".docx"
        return FileResponse(
            temp_out,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=out_name
        )
    except Exception as e:
        cleanup_file(temp_in)
        cleanup_file(temp_out)
        raise HTTPException(status_code=500, detail=f"Failed to convert PDF to Word: {str(e)}")

# -------------------------------------------------------------
# 7. PDF TO POWERPOINT (.pdf -> .pptx)
# -------------------------------------------------------------
@app.post("/api/pdf-to-powerpoint")
async def pdf_to_powerpoint(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...)
):
    temp_in = os.path.join(TEMP_DIR, f"in_{uuid.uuid4().hex}.pdf")
    temp_out = os.path.join(TEMP_DIR, f"out_{uuid.uuid4().hex}.pptx")

    try:
        with open(temp_in, "wb") as f:
            content = await file.read()
            f.write(content)

        reader = pypdf.PdfReader(temp_in)
        prs = Presentation()
        # Set 16:9 or standard slide dimensions
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)
        blank_slide_layout = prs.slide_layouts[6]

        for idx, page in enumerate(reader.pages):
            slide = prs.slides.add_slide(blank_slide_layout)
            
            # Add slide header
            tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.4), Inches(0.8))
            tf = tx_box.text_frame
            p = tf.paragraphs[0]
            p.text = f"Slide {idx + 1}"
            p.font.size = Pt(20)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)

            # Extract page text
            text = page.extract_text()
            if text:
                content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.4), Inches(5.2))
                ctf = content_box.text_frame
                ctf.word_wrap = True
                
                lines = [l.strip() for l in text.splitlines() if l.strip()]
                for l_idx, line in enumerate(lines[:12]):  # fit top text lines nicely
                    cp = ctf.paragraphs[0] if l_idx == 0 else ctf.add_paragraph()
                    cp.text = f"• {line}"
                    cp.font.size = Pt(13)
                    cp.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            else:
                content_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(8.4), Inches(2.0))
                ctf = content_box.text_frame
                cp = ctf.paragraphs[0]
                cp.text = f"[Page {idx + 1} Visual Content]"
                cp.font.size = Pt(14)
                cp.font.italic = True

        prs.save(temp_out)

        background_tasks.add_task(cleanup_file, temp_in)
        background_tasks.add_task(cleanup_file, temp_out)

        out_name = (file.filename or "presentation.pdf").rsplit(".", 1)[0] + ".pptx"
        return FileResponse(
            temp_out,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            filename=out_name
        )
    except Exception as e:
        cleanup_file(temp_in)
        cleanup_file(temp_out)
        raise HTTPException(status_code=500, detail=f"Failed to convert PDF to PowerPoint: {str(e)}")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)
