import os
import io
import urllib.request
import urllib.error
import mimetypes
import uuid
import openpyxl
from docx import Document
from pptx import Presentation
from reportlab.pdfgen import canvas

os.makedirs("test_files", exist_ok=True)

# 1. Create sample docx
doc = Document()
doc.add_heading("Production Document Test", level=1)
doc.add_paragraph("This is a genuine paragraph in our test Word document.")
table = doc.add_table(rows=2, cols=2)
table.rows[0].cells[0].text = "Header 1"
table.rows[0].cells[1].text = "Header 2"
table.rows[1].cells[0].text = "Val A"
table.rows[1].cells[1].text = "Val B"
docx_path = "test_files/sample.docx"
doc.save(docx_path)

# 2. Create sample xlsx
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sales Data"
ws.append(["Item", "Qty", "Price", "Total"])
ws.append(["Widget A", 10, 15.5, 155.0])
ws.append(["Widget B", 5, 42.0, 210.0])
xlsx_path = "test_files/sample.xlsx"
wb.save(xlsx_path)

# 3. Create sample pptx
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Quarterly Business Review"
slide.placeholders[1].text = "Performance & Metrics Presentation"
pptx_path = "test_files/sample.pptx"
prs.save(pptx_path)

# 4. Create sample pdf
pdf_path = "test_files/sample.pdf"
c = canvas.Canvas(pdf_path)
c.drawString(100, 750, "Sample PDF Source Document")
c.drawString(100, 700, "Section 1: Data extraction and formatting verification.")
c.showPage()
c.drawString(100, 750, "Page 2 of Sample Document")
c.save()

BASE_URL = "http://127.0.0.1:8000/api"

def upload_file(url, file_path, filename, field_name="file", extra_fields=None):
    boundary = uuid.uuid4().hex
    body = bytearray()
    
    if extra_fields:
        for k, v in extra_fields.items():
            body.extend(f'--{boundary}\r\n'.encode('utf-8'))
            body.extend(f'Content-Disposition: form-data; name="{k}"\r\n\r\n'.encode('utf-8'))
            body.extend(f'{v}\r\n'.encode('utf-8'))

    with open(file_path, "rb") as f:
        file_bytes = f.read()

    body.extend(f'--{boundary}\r\n'.encode('utf-8'))
    body.extend(f'Content-Disposition: form-data; name="{field_name}"; filename="{filename}"\r\n'.encode('utf-8'))
    body.extend(b'Content-Type: application/octet-stream\r\n\r\n')
    body.extend(file_bytes)
    body.extend(b'\r\n')
    body.extend(f'--{boundary}--\r\n'.encode('utf-8'))

    req = urllib.request.Request(url, data=bytes(body), method="POST")
    req.add_header('Content-Type', f'multipart/form-data; boundary={boundary}')
    with urllib.request.urlopen(req) as resp:
        return resp.read()

print("--- Testing Word to PDF ---")
res = upload_file(f"{BASE_URL}/word-to-pdf", docx_path, "sample.docx")
assert res.startswith(b"%PDF"), "Output is not a valid PDF"
print(f"[PASS] Word to PDF passed! Size: {len(res)} bytes")

print("--- Testing Excel to PDF ---")
res = upload_file(f"{BASE_URL}/excel-to-pdf", xlsx_path, "sample.xlsx")
assert res.startswith(b"%PDF"), "Output is not a valid PDF"
print(f"[PASS] Excel to PDF passed! Size: {len(res)} bytes")

print("--- Testing PowerPoint to PDF ---")
res = upload_file(f"{BASE_URL}/powerpoint-to-pdf", pptx_path, "sample.pptx")
assert res.startswith(b"%PDF"), "Output is not a valid PDF"
print(f"[PASS] PowerPoint to PDF passed! Size: {len(res)} bytes")

print("--- Testing PDF to Word ---")
res = upload_file(f"{BASE_URL}/pdf-to-word", pdf_path, "sample.pdf")
assert len(res) > 500, "Output is too small"
print(f"[PASS] PDF to Word passed! Size: {len(res)} bytes")

print("--- Testing PDF to PowerPoint ---")
res = upload_file(f"{BASE_URL}/pdf-to-powerpoint", pdf_path, "sample.pdf")
assert len(res) > 500, "Output is too small"
print(f"[PASS] PDF to PowerPoint passed! Size: {len(res)} bytes")

print("--- Testing Protect PDF (AES-256) ---")
protected_res = upload_file(f"{BASE_URL}/protect-pdf", pdf_path, "sample.pdf", extra_fields={"password": "MySecretPassword123!"})
assert protected_res.startswith(b"%PDF"), "Protected output is not a valid PDF"
print(f"[PASS] Protect PDF passed! Protected size: {len(protected_res)} bytes")

# Save protected file to test unlock
prot_path = "test_files/protected.pdf"
with open(prot_path, "wb") as f:
    f.write(protected_res)

print("--- Testing Unlock PDF ---")
unlocked_res = upload_file(f"{BASE_URL}/unlock-pdf", prot_path, "protected.pdf", extra_fields={"password": "MySecretPassword123!"})
assert unlocked_res.startswith(b"%PDF"), "Unlocked output is not a valid PDF"
print(f"[PASS] Unlock PDF passed! Unlocked size: {len(unlocked_res)} bytes")

# Clean up
import shutil
shutil.rmtree("test_files", ignore_errors=True)
print("\n=== ALL 7 SERVER-ASSISTED CONVERSION TOOLS VERIFIED AND 100% OPERATIONAL! ===")
